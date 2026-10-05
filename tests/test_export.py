"""Выгрузка плана: заполненные колонки датасета, листы по командам, PDF."""

import io
from pathlib import Path

import openpyxl
import pytest

from backend.export.report import NOT_IN_QUARTER, excel_bytes, pdf_bytes, plan_rows
from backend.recalculation.pipeline import week0_from_excel

DATASET = Path("data/raw/Хакатон_датасетс_правками_по_списку_вопросов_1.xlsx")


@pytest.fixture(scope="module")
def state():
    return week0_from_excel(DATASET.read_bytes())


def test_plan_rows_fill_dataset_columns(state):
    """Колонки, пустые в датасете у ToDo, заполняет планировщик."""
    rows = plan_rows(state["plan"])
    included = [r for r in rows if r["Решение"] == "Включено в квартал"]
    assert included
    for row in included:
        assert row["будет включено в спринт"].startswith("Спринт")
        assert row["плановый результат"], row["task_id"]
        assert row["исполнитель"], row["task_id"]

    postponed = [r for r in rows if r["Решение"] == "Перенесено"]
    assert all(r["будет включено в спринт"] == NOT_IN_QUARTER for r in postponed)
    assert all(r["Причина"] for r in postponed)


def test_plan_rows_team_filter(state):
    rows = plan_rows(state["plan"], team="Team-L")
    assert rows and {r["team_id"] for r in rows} == {"Team-L"}


def test_excel_has_sheet_per_team(state):
    data = excel_bytes(state["plan"], state["alerts"], state["starmap"])
    book = openpyxl.load_workbook(io.BytesIO(data))
    assert {"Сводка", "План", "Алерты"} <= set(book.sheetnames)
    teams = {t["team_id"] for t in state["plan"]["current"] if t.get("team_id")}
    assert teams <= set(book.sheetnames)
    plan_sheet = book["План"]
    assert plan_sheet.max_row > 1
    header = [c.value for c in plan_sheet[1]]
    for column in ("плановый результат", "будет включено в спринт", "исполнитель"):
        assert column in header


def test_pdf_is_valid_and_not_empty(state):
    data = pdf_bytes(state["plan"], state["alerts"], state["starmap"], team="Team-L")
    assert data.startswith(b"%PDF")
    assert len(data) > 5000


def test_fact_template_round_trip(state):
    """Шаблон факта должен приниматься обратно как факт спринта."""
    import copy

    from backend.export.report import fact_template_bytes, fact_template_rows
    from backend.recalculation.pipeline import fact_from_any_excel
    from backend.recalculation.recalc import apply_sprint_fact

    plan = copy.deepcopy(state["plan"])
    engineers = state["parsed"]["engineers_profiles"]
    rows = fact_template_rows(plan, engineers, 1)
    assert rows and all(r["task_id"] and r["status"] for r in rows)

    data = fact_template_bytes(plan, engineers, 1)
    fact = fact_from_any_excel(data, plan)
    assert set(fact["tasks"]) == {r["task_id"] for r in rows}

    updated = apply_sprint_fact(plan=plan, fact=fact, dependencies=state["clean_deps"],
                                engineers=engineers, capacity=plan["capacity"])
    assert updated["current_sprint"] == 1
    # Работа ровно по плану: срывов и каскадных сдвигов быть не должно
    types = {a.type if hasattr(a, "type") else a["type"] for a in updated.get("alerts") or []}
    assert not types & {"red", "yellow"}
